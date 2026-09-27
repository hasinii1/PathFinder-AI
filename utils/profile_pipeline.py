import json
import os
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


# ============================================================
# OUTPUT FILES
# ============================================================

STUDENT_DATA_DIR = (
    ROOT
    / "data"
    / "students"
)

PROFILE_FILE = (
    STUDENT_DATA_DIR
    / "profile_output.json"
)

ASSESSMENT_FILE = (
    STUDENT_DATA_DIR
    / "assessment_output.json"
)

SKILL_GAP_FILE = (
    STUDENT_DATA_DIR
    / "skill_gap_output.json"
)

OPPORTUNITY_FILE = (
    STUDENT_DATA_DIR
    / "opportunity_results.json"
)

ROADMAP_FILE = (
    STUDENT_DATA_DIR
    / "roadmap_output.json"
)


# ============================================================
# CLEAR OLD STUDENT RESULTS
# ============================================================

def clear_previous_results():
    """
    Remove previous student's generated results.

    This prevents one student's information from appearing
    when another student's resume is uploaded.
    """

    files_to_clear = [
        PROFILE_FILE,
        ASSESSMENT_FILE,
        SKILL_GAP_FILE,
        OPPORTUNITY_FILE,
        ROADMAP_FILE,
    ]

    STUDENT_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    for file_path in files_to_clear:

        try:

            if file_path.exists():
                file_path.unlink()

        except OSError:
            # Do not stop the application if a file is
            # temporarily locked.
            pass


# ============================================================
# EXTRACT RESUME TEXT
# ============================================================

def extract_resume_text(
    file_bytes: bytes,
    filename: str,
) -> str:

    suffix = Path(filename).suffix.lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if suffix == ".pdf":

        import pymupdf

        document = pymupdf.open(
            stream=file_bytes,
            filetype="pdf",
        )

        parts = []

        for page in document:

            text = page.get_text(
                "text"
            )

            if text and text.strip():
                parts.append(
                    text.strip()
                )

        document.close()

        return "\n\n".join(parts)

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if suffix == ".docx":

        from docx import Document
        import io

        document = Document(
            io.BytesIO(file_bytes)
        )

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:

            for row in table.rows:

                row_text = " ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )

                if row_text:
                    parts.append(row_text)

        return "\n".join(parts)

    raise ValueError(
        "Please upload a PDF or DOCX resume."
    )


# ============================================================
# SAVE PROFILE
# ============================================================

def save_profile(
    profile: dict,
):
    """Save the current student's profile."""

    STUDENT_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    PROFILE_FILE.write_text(
        json.dumps(
            profile,
            indent=4,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    return PROFILE_FILE


# ============================================================
# RUN COMPLETE PIPELINE
# ============================================================

def run_pipeline(
    profile: dict,
):
    """
    Run Members 2–4 and opportunity matching
    using the newly analyzed profile.
    """

    # --------------------------------------------------------
    # Save the NEW profile first
    # --------------------------------------------------------

    save_profile(profile)

    # --------------------------------------------------------
    # MEMBER 2 - SKILL ASSESSMENT
    # --------------------------------------------------------

    from ml.skill_assessment.skill_assessment import (
        build_assessment,
        save_output as save_assessment,
    )

    assessment = build_assessment(
        profile
    )

    save_assessment(
        profile,
        assessment,
    )

    # --------------------------------------------------------
    # MEMBER 2 - SKILL GAP
    # --------------------------------------------------------

    from ml.skill_gap.skill_gap import (
        calculate_skill_gaps,
        load_careers,
        save_output as save_gap,
    )

    careers_df = load_careers()

    gap_analysis = calculate_skill_gaps(
        profile,
        {
            "skill_assessment": assessment
        },
        careers_df,
    )

    save_gap(
        profile,
        gap_analysis,
    )

    # --------------------------------------------------------
    # MEMBER 4 - ROADMAP
    # --------------------------------------------------------

    from ml.roadmap.roadmap_utils import (
        run_member_4_pipeline,
    )

    old_cwd = os.getcwd()

    os.chdir(ROOT)

    try:

        run_member_4_pipeline()

    finally:

        os.chdir(old_cwd)

    # --------------------------------------------------------
    # MEMBER 3 - OPPORTUNITY MATCHING
    # --------------------------------------------------------

    from ml.opportunity_matching.opportunity_matching import (
        load_skill_gap,
        load_opportunities,
        match_opportunities,
        OUTPUT_FILE,
    )

    skill_gap_data = load_skill_gap()

    opportunities = load_opportunities()

    result = match_opportunities(
        skill_gap_data,
        opportunities,
    )

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT_FILE.write_text(
        json.dumps(
            result,
            indent=4,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    # --------------------------------------------------------
    # RETURN EVERYTHING
    # --------------------------------------------------------

    return {
        "profile": profile,
        "assessment": assessment,
        "skill_gap": gap_analysis,
        "opportunities": result,
    }